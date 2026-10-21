import statistics as st, collections
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PS=[2,4,5,6,7,8,10,12,14]; PS7=[2,4,6,8,10,12,14]
raw=[]
for l in open('data.txt'):
    if l.startswith('#') or not l.strip(): continue
    t=l.split(); im=t[0]; rep=int(t[1]); v=list(map(float,t[2:]))
    raw.append((im,'seq',1,rep,v[0],v[1]))
    for i,p in enumerate(PS):
        e=v[2+4*i:6+4*i]; raw.append((im,'equal',p,rep,e[0],e[1])); raw.append((im,'adaptive',p,rep,e[2],e[3]))
D=collections.defaultdict(list)
for im,s,p,rep,L,T in raw: D[(im,s,p)].append((L,T))
med=lambda im,s,p,k: st.median(x[k] for x in D[(im,s,p)])
sd=lambda im,s,p,k: st.stdev(x[k] for x in D[(im,s,p)])
auto=lambda p:'equal' if p<=6 else 'adaptive'
S=lambda im,p,k,s=None: med(im,'seq',1,k)/med(im,s or auto(p),p,k)
mean=lambda a:sum(a)/len(a)
KF=lambda s,p:(1/s-1/p)/(1-1/p)
ORD=['lena_gray','boat_gray','cameraman_gray','peppers_gray','lena_color','peppers_color2','f16_color','goldhill_color']
NM={'lena_gray':'Lena (gray)','boat_gray':'Boat (gray)','cameraman_gray':'Cameraman (gray)','peppers_gray':'Peppers (gray)','lena_color':'Lena (color)','peppers_color2':'Peppers (color)','f16_color':'F16 (color)','goldhill_color':'Gold Hill (color)'}
G=ORD[:4]; C=ORD[4:]; ALL=ORD

doc=Document()
sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(sec,m,Cm(1.8))
stl=doc.styles['Normal']; stl.font.name='Times New Roman'; stl.font.size=Pt(11)
stl.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')

def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),color); tcPr.append(sh)
def para(text,bold=False,italic=False,size=11,align=None,space_after=4):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size)
    p.paragraph_format.space_after=Pt(space_after)
    if align: p.alignment=align
    return p
def table(caption,header,rows,note=None,bold_rows=(),fs=9.5):
    cp=para(caption,bold=True,size=11,space_after=3); cp.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(header)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(header):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(fs); c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; shade(c,'D9E2F3')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=''; r=cells[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(fs); r.bold = ri in bold_rows
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
            if ri in bold_rows: shade(cells[i],'F2F2F2')
    if note: para(note,italic=True,size=9.5,space_after=10)
    else: para('',space_after=6)

para('Experimental results: tables from one measurement session',bold=True,size=16,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=6)
para('Platform: Intel Core i7-13650HX (6 P-cores + 8 E-cores), Ubuntu 24.04, GCC 13.3 (-O2), Open MPI 4.1.6. Protocol: AC power, governor = performance, MPI processes bound to cores, sequential baseline pinned to one P-core; median of 5 runs (repetitions as the outer loop). Programs: codes_min.c (sequential) and codep_min.c (parallel). Distribution rule: Equal for P ≤ 6, Adaptive for P > 6, unless stated otherwise.',size=10,space_after=4)
para('LBG time: codebook training only. Total processing time: all processing stages (building and converting the training vectors, [parallel only: work distribution, speed calibration, scatter, gather], LBG, 8-bit codebook, delta coding, Huffman container, decoding); it excludes image reading, console output, file writing, entropy statistics, and the quality metrics (MSE, PSNR, SSIM).',size=10,space_after=4)
para('S = T_seq / T_P (per image, then the arithmetic mean over the images); E = S / P; Karp–Flatt e = (1/S − 1/P) / (1 − 1/P), computed from the mean S. Efficiency and Karp–Flatt are computed from unrounded speedups.',size=10,space_after=10)

n=1
table(f'Table {n}. Sequential execution times (median of 5 runs, s).',['Image','LBG time','Total time','Outside LBG (ms)','LBG share (%)','SD of Total (ms)'],
 [[NM[i],f"{med(i,'seq',1,0):.4f}",f"{med(i,'seq',1,1):.4f}",f"{1000*(med(i,'seq',1,1)-med(i,'seq',1,0)):.1f}",f"{100*med(i,'seq',1,0)/med(i,'seq',1,1):.2f}",f"{1000*sd(i,'seq',1,1):.2f}"] for i in ORD]); n+=1
for k,name in ((0,'LBG time'),(1,'Total processing time')):
    table(f'Table {n}. {name} (s, median of 5 runs).',['Image','Seq']+[f'P = {p}' for p in PS7],
      [[NM[i],f"{med(i,'seq',1,k):.4f}"]+[f"{med(i,auto(p),p,k):.4f}" for p in PS7] for i in ORD]); n+=1
for k,name in ((0,'LBG'),(1,'Total processing')):
    table(f'Table {n}. {name} speedup of each image.',['Image']+[f'P = {p}' for p in PS7],
      [[NM[i]]+[f"{S(i,p,k):.2f}" for p in PS7] for i in ORD]+[['Mean (8 images)']+[f"{mean([S(i,p,k) for i in ALL]):.2f}" for p in PS7]],bold_rows=(8,)); n+=1
rows=[]
for p in PS:
    sl=mean([S(i,p,0) for i in ALL]); sg=mean([S(i,p,1) for i in G]); sc=mean([S(i,p,1) for i in C]); sa=mean([S(i,p,1) for i in ALL])
    rows.append([p,f"{sl:.2f}",f"{100*sl/p:.1f}",f"{KF(sl,p):.3f}",f"{sg:.2f}",f"{sc:.2f}",f"{sa:.2f}",f"{100*sg/p:.1f}",f"{100*sc/p:.1f}",f"{KF(sg,p):.3f}",f"{KF(sc,p):.3f}"])
table(f'Table {n}. Speedup, parallel efficiency and Karp–Flatt metric (means).',['P','LBG S','LBG E (%)','LBG KF','Total S gray','Total S color','Total S mean','E gray (%)','E color (%)','KF gray','KF color'],rows,
      note='LBG: mean of the 8 images; Total: mean of the 4 grayscale and of the 4 color images (and of all 8).',fs=9); n+=1
rows=[]
for p in PS7:
    rows.append([p]+[f"{mean([1000*(med(i,auto(p),p,1)-med(i,auto(p),p,0)) for i in g]):.2f}" for g in (G,C)]+[f"{mean([100*med(i,auto(p),p,0)/med(i,auto(p),p,1) for i in g]):.1f}" for g in (G,C)])
table(f'Table {n}. Time outside LBG and the share of LBG in the Total processing time.',['P','Outside LBG, gray (ms)','Outside LBG, color (ms)','LBG share, gray (%)','LBG share, color (%)'],rows); n+=1
rows=[]
for p in PS:
    e=mean([S(i,p,0,'equal') for i in ALL]); a=mean([S(i,p,0,'adaptive') for i in ALL]); te=mean([S(i,p,1,'equal') for i in ALL]); ta=mean([S(i,p,1,'adaptive') for i in ALL])
    rows.append([p,f"{e:.2f}",f"{a:.2f}",f"{a-e:.2f}",f"{te:.2f}",f"{ta:.2f}",f"{ta-te:.2f}"])
table(f'Table {n}. Equal versus Adaptive workload distribution (mean speedup of the 8 images, same session).',['P','LBG Equal','LBG Adaptive','LBG difference','Total Equal','Total Adaptive','Total difference'],rows,
      note='Difference = S_Adaptive − S_Equal (computed from unrounded means).'); n+=1
rows=[]
for p in (8,10,12,14):
    lr=[100*(1-med(i,'adaptive',p,0)/med(i,'equal',p,0)) for i in ALL]; tr=[100*(1-med(i,'adaptive',p,1)/med(i,'equal',p,1)) for i in ALL]
    rows.append([p,f"{mean(lr):.1f}",f"{min(lr):.1f}–{max(lr):.1f}",f"{mean(tr):.1f}",f"{min(tr):.1f}–{max(tr):.1f}"])
table(f'Table {n}. Time reduction of Adaptive compared with Equal distribution (%).',['P','LBG time (mean)','LBG time (min–max)','Total time (mean)','Total time (min–max)'],rows,
      note='Reduction (%) = (1 − T_Adaptive / T_Equal) × 100, per image, then the mean over the 8 images.'); n+=1
table(f'Table {n}. Run-to-run variation of the Total processing time at P = 14 (Adaptive, 5 runs).',['Image','Median (s)','SD (ms)','SD (%)','Min–max (s)'],
 [[NM[i],f"{med(i,'adaptive',14,1):.4f}",f"{1000*sd(i,'adaptive',14,1):.2f}",f"{100*sd(i,'adaptive',14,1)/med(i,'adaptive',14,1):.2f}",f"{min(x[1] for x in D[(i,'adaptive',14)]):.4f}–{max(x[1] for x in D[(i,'adaptive',14)]):.4f}"] for i in ORD],
 note='The reported times are medians, which are not affected by the occasional slow run (for example, Peppers (color)).'); n+=1
comp=[('Lena','Gray','23.99','0.333','70.13','29.67','0.8190'),('Boat','Gray','27.18','0.294','79.88','29.11','0.8394'),('Cameraman','Gray','29.60','0.270','68.84','29.75','0.8855'),('Peppers','Gray','25.30','0.316','74.47','29.41','0.7624'),('Mean (gray)','Gray','26.52','0.304','73.33','29.49','0.8266'),
      ('Lena','Color','24.52','0.979','83.60','28.91','0.7594'),('Peppers','Color','24.40','0.984','95.00','28.35','0.6964'),('F16','Color','27.07','0.887','73.41','29.47','0.8539'),('Gold Hill','Color','24.03','0.999','92.97','28.45','0.7520'),('Mean (color)','Color','25.00','0.962','86.25','28.80','0.7654')]
table(f'Table {n}. Compression results (4 × 4 blocks, K = 64).',['Image','Type','CR','BPP','MSE','PSNR (dB)','SSIM'],[list(r) for r in comp],bold_rows=(4,9),
      note='CR = original size / size of the compressed container (14-byte header + code table + codebook + Huffman-coded data). Color images: MSE and PSNR over the RGB channels; SSIM is the mean of the channel-wise values. Means are means of the per-image values.')
doc.save('Experimental_Tables_RUN1.docx'); print("tables:",n)
