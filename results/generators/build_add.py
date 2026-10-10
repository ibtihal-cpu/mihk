import csv, statistics as st, openpyxl
from decimal import Decimal as Dm, ROUND_HALF_UP
def q(vals,d):
    m=sum(Dm(str(v)) for v in vals)/len(vals)
    return str(m.quantize(Dm(1).scaleb(-d),rounding=ROUND_HALF_UP))
from docx import Document
from docx.shared import Pt
R='/home/user/mihk/results/'
SP='/tmp/claude-0/-home-user-mihk/4ae9f12a-c6d9-593e-bdd1-593721fcd2b9/scratchpad/ps/ps.csv'
t2="""0.4406 0.5679 0.4900 0.2577 0.2554 0.3251 0.4465 0.5589
0.2653 0.2326 0.2372 0.2798 0.2324 0.2093 0.2895 0.2446
0.1572 0.1674 0.1485 0.1923 0.2071 0.1848 0.1437 0.1354
0.1166 0.1096 0.1327 0.0967 0.1138 0.0564 0.1619 0.1046
0.6252 0.5190 0.5809 0.5386 0.6054 0.4529 0.4837 0.4413""".split('\n')
names=['Lena','Boat','Cameraman','Peppers','Chest_X-ray','Mammogram','Lung_CT','Brain_MRI']
vs=["2x2/K64","4x4/K64","8x8/K64","4x4/K32","4x4/K128"]
rows=list(csv.DictReader(open(SP)))
ps=[]
for i,r in enumerate(rows):
    v=i//8; j=i%8
    ps.append(dict(variant=r['variant'],block=int(r['block']),K=int(r['K']),image=names[j],PSNR=float(r['PSNR']),SSIM=float(r['SSIM']),CR=float(r['CR']),iters=int(r['it']),time=float(t2[v].split()[j])))
    assert r['variant']==vs[v]
with open(R+'data/param_study.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['variant','block','K','image','PSNR_dB','SSIM','system_CR','LBG_iters','LBG_time_median_s'])
    for p in ps: w.writerow([p['variant'],p['block'],p['K'],p['image'],f"{p['PSNR']:.2f}",f"{p['SSIM']:.4f}",f"{p['CR']:.2f}",p['iters'],f"{p['time']:.4f}"])
def vals(v,k): return [p[k] for p in ps if p['variant']==v]
def prow(v,label): return [label,q(vals(v,'PSNR'),2),q(vals(v,'SSIM'),4),q(vals(v,'CR'),2),q(vals(v,'iters'),1),q(vals(v,'time'),4)]
H=['PSNR (dB)','SSIM','CR','LBG iterations','LBG time (s)']
tabs=[]
tabs.append(('Table 4.A. Effect of the block size (K = 64; mean of the eight grayscale images).',['Block size']+H,[prow('2x2/K64','2×2'),prow('4x4/K64','4×4'),prow('8x8/K64','8×8')]))
tabs.append(('Table 4.B. Effect of the codebook size (4×4 blocks; mean of the eight grayscale images).',['K']+H,[prow('4x4/K32','32'),prow('4x4/K64','64'),prow('4x4/K128','128')]))
# entropy coding
wb=openpyxl.load_workbook(R+'final_results.xlsx',data_only=True); ws=wb['Entropy']
er=list(ws.iter_rows(values_only=True))[1:]
cb=64*16; N=16384; O=262144
cr8=O/(cb+N); cr6=O/(cb+N*6/8)
erows=[];ent=[]
for r in er:
    nm,Hi,Hd,_,_,_,_,_,cr=r
    gain=(cr/cr6-1)*100
    erows.append([nm.replace(' ','_') if False else nm,f"{cr8:.2f}",f"{cr6:.2f}",f"{cr:.2f}",f"{gain:+.1f}"])
    ent.append((nm,cr8,cr6,cr,gain))
erows.append(['Mean',f"{cr8:.2f}",f"{cr6:.2f}",q([e[3] for e in ent],2),f"{st.mean(e[4] for e in ent):+.1f}"])
tabs.append(('Table 4.C. Effect of index coding on the system compression ratio (4×4 blocks, K = 64; codebook of 1024 bytes included in all cases).',['Image','Fixed 8-bit index','Fixed 6-bit index','Delta + Huffman (proposed)','Gain over 6-bit (%)'],erows))
# baselines
BL=[l for l in csv.DictReader(open(R+'data/baselines.csv'))]
meth=['Proposed','JPEG','JPEG 2000']
def bm(m,k): return st.mean(float(b[k]) for b in BL if b['method']==m)
tabs.append(('Table 4.D. Rate-matched reference comparison with JPEG and JPEG 2000 (mean of the eight grayscale images).',['Method','CR','PSNR (dB)','SSIM'],
 [[{'Proposed':'Proposed (LBG + Delta + Huffman + BF)'}.get(m,m),q([b['CR'] for b in BL if b['method']==m],2),q([b['PSNR_dB'] for b in BL if b['method']==m],2),q([b['SSIM'] for b in BL if b['method']==m],4)] for m in meth]))
per=[]
for n in names:
    b={x['method']:x for x in BL if x['image']==n}
    per.append([n]+sum([[b[m]['CR'],b[m]['PSNR_dB'],b[m]['SSIM']] for m in meth],[]))
tabs.append(('Table A.x. Per-image reference comparison with JPEG and JPEG 2000 (CR / PSNR in dB / SSIM).',['Image','Prop. CR','Prop. PSNR','Prop. SSIM','JPEG CR','JPEG PSNR','JPEG SSIM','J2K CR','J2K PSNR','J2K SSIM'],per))
doc=Document(); doc.styles['Normal'].font.name='Times New Roman'; doc.styles['Normal'].font.size=Pt(10)
doc.add_heading('Chapter 4 additional tables (parameter study, index coding, JPEG/JPEG 2000 reference)',1)
for cap,hdr,rws in tabs:
    p=doc.add_paragraph(); p.add_run(cap).bold=True
    t=doc.add_table(rows=1,cols=len(hdr)); t.style='Table Grid'
    for i,h in enumerate(hdr): t.rows[0].cells[i].text=h
    for r in rws:
        c=t.add_row().cells
        for i,x in enumerate(r): c[i].text=str(x)
    doc.add_paragraph()
doc.save(R+'CH4_ADDITIONS.docx')
# xlsx
wb=openpyxl.load_workbook(R+'final_results.xlsx')
for n in ['Param study','Baselines JPEG-J2K','Index coding']:
    if n in wb.sheetnames: del wb[n]
w=wb.create_sheet('Param study'); w.append(['variant','block','K','image','PSNR_dB','SSIM','system_CR','LBG_iters','LBG_time_median_s'])
for p in ps: w.append([p['variant'],p['block'],p['K'],p['image'],p['PSNR'],p['SSIM'],p['CR'],p['iters'],p['time']])
w=wb.create_sheet('Baselines JPEG-J2K'); w.append(['image','method','setting','CR','PSNR_dB','SSIM'])
for b in BL: w.append([b['image'],b['method'],b['setting'],float(b['CR']),float(b['PSNR_dB']),float(b['SSIM'])])
w=wb.create_sheet('Index coding'); w.append(['image','CR_fixed8','CR_fixed6','CR_delta_huffman','gain_over_6bit_pct'])
for e in ent: w.append([e[0],round(e[1],2),round(e[2],2),e[3],round(e[4],1)])
wb.save(R+'final_results.xlsx')
for cap,hdr,rws in tabs[:4]:
    print(cap); print(hdr)
    for r in rws: print(r)
