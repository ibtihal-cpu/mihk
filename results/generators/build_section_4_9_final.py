#!/usr/bin/env python3
"""Final Tables 4.16 / 4.17 + discussion text (contextual literature comparison). Output: results/Section_4_9_final.docx
All 'This work' numbers come from results/final_campaign (correctness.csv, timing_final.csv)."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
for s in d.sections: s.left_margin = s.right_margin = Cm(2.5)

def para(t, bold=False, italic=False, size=None, align=None, space=6):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space); return p

def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "bottom", "insideH"):
        x = OxmlElement("w:" + e); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "6"); x.set(qn("w:color"), "000000"); b.append(x)
    tblPr.append(b)

def table(head, rows, widths, bold_rows=(), size=10):
    t = d.add_table(rows=1, cols=len(head)); borders(t)
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(size); r.bold = ri in bold_rows
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t

# ---------------- Table 4.16 ----------------
para("Table 4.16. Compression ratio, PSNR and time ratio reported in selected studies and obtained in this work for standard test images (contextual comparison).", bold=True, size=11)
rows = [
 ["[6]", "Lena", "Thomas et al.", "1.52", "54.28", "–"],
 ["[16]", "", "Ranjan and Kumar", "4.41", "34.78", "–"],
 ["", "", "This work", "23.50", "29.67", "9.21×"],
 ["[6]", "Boat", "Thomas et al.", "1.41", "54.07", "–"],
 ["", "", "This work", "26.54", "29.11", "9.22×"],
 ["[17]", "Cameraman", "Bilal et al.", "–", "25.3–26.8*", "2.61×"],
 ["", "", "This work", "28.88", "29.75", "9.10×"],
]
table(["Ref.", "Image", "Reference", "CR", "PSNR (dB)", "Time ratio"], rows, [1.3, 2.4, 4.2, 1.8, 2.6, 2.4], bold_rows=(2, 4, 6))
para("Thomas et al.: integer DWT + coefficient reduction + Huffman coding (Tables 1 and 3); Ranjan and Kumar: 2D DWT + PCA + canonical Huffman coding (Table 1, 512×512); Bilal et al.: FLBG; this work: LBG + Delta + Huffman coding (sequential, CR as system ratio). * Bilal et al. report PSNR versus bit rate (≤ 0.625 bpp) and no CR. “–” = not reported. Time ratio is computed within each study: LBG time divided by FLBG time in Bilal et al. (Cameraman, codebook size 16, 0.25 bpp, Table 6), and sequential LBG time divided by MPI LBG time at P = 14 in this work.", italic=True, size=9.5)

para("Table 4.16 presents the compression ratio (CR) and PSNR reported in selected image-compression studies alongside the corresponding results of the proposed framework for the common test images. The reported studies include wavelet-based methods (Thomas et al. [6]; Ranjan and Kumar [16]) and a fast LBG variant (Bilal et al. [17]). The wavelet-based methods operate at low compression ratios (1.41–4.41:1) with high PSNR (34.78–54.28 dB), whereas the proposed framework operates at 23.50–28.88:1 with PSNR between 29.11 and 29.75 dB. The studies therefore occupy different regions of the rate–distortion trade-off, and the table is not intended to establish the superiority of one method over another. Bilal et al. report PSNR as a function of bit rate; their values (25.3–26.8 dB at up to 0.625 bpp for Cameraman) are given for context only. The last column reports the time ratio computed within each study: the LBG-to-FLBG time ratio in Bilal et al. (codebook size 16, 0.25 bpp), and the sequential-to-parallel LBG time ratio at P = 14 in this work. These ratios reflect different mechanisms (algorithmic complexity reduction versus parallel execution) and different hardware, and absolute times are not compared. Differences in test-image files, implementations, and compressed-size accounting further limit direct comparison; the values are therefore indicative only.")
para("Parallel performance was also compared with selected studies that applied parallel computing to clustering or workload distribution. Table 4.17 summarizes the computing resources and the reported speedup of each study. The reported speedups are relative to different baselines, problem sizes, and hardware, so the figures are used to place the obtained speedup in context rather than to rank the implementations.")

# ---------------- Table 4.17 ----------------
para("Table 4.17. Contextual comparison of reported parallel speedups.", bold=True, size=11)
rows2 = [
 ["Tanović et al. (2024)", "Parallel k-means for image classification (MNIST variants)", "12 processor cores, OpenMP and MPI", "Not stated in the abstract", "Up to 10.2×"],
 ["Papakyriakou et al. (2025)", "MPI-based k-means clustering of 75 million synthetic 2D points", "2–48 MPI processes on a Raspberry Pi cluster", "Two-process run", "≈10.05× (181.83 s to 18.09 s); efficiency 41.89% at 48 processesᵃ"],
 ["Bystrov et al. (2023)", "Discrete-element simulation with resource-aware weighted RCB partitioning", "Eight heterogeneous containers, OpenCL and MPI", "Not stated in the abstract", "Up to 6.0×"],
 ["This work", "Parallel LBG-based image compression", "14 MPI processes mapped to 6 P-cores and 8 E-cores", "Sequential program", "9.19× average LBG-stage speedup; 5.00× (grayscale) and 6.84× (color) total speedup at P = 14"],
]
table(["Study", "Application", "Computing resources", "Baseline", "Reported speedup"], rows2, [2.9, 4.0, 3.3, 2.4, 3.8], bold_rows=(3,), size=9.5)
para("ᵃ Papakyriakou et al. used the two-process execution as the baseline (no sequential run); the speedup is computed from the total times as T2/T48, and the efficiency is relative to the same baseline. The studies differ in application, workload, hardware, process count, and speedup baseline; therefore, the values represent a contextual comparison rather than a direct performance ranking. The 9.19× value is the mean LBG-stage speedup over the eight test images; the total speedup includes the non-parallel stages.", italic=True, size=9.5)

para("Bystrov et al. applied resource-aware weighted recursive coordinate bisection to distribute a discrete-element simulation across heterogeneous cloud resources. On seven heterogeneous containers, the weighted strategy reduced execution time by up to 32.6% compared with unweighted repartitioning. The implementation also achieved a maximum speedup of 6.0× when executed on eight heterogeneous containers. Although the application differs from image compression, its resource-aware partitioning is conceptually related to the adaptive workload-distribution strategy used in the present work, which reduced the LBG time by 17.2–28.3% relative to equal distribution at P = 8–14.")
d.save("../Section_4_9_final.docx"); print("saved results/Section_4_9_final.docx")
